"""
Integration test dell'export Excel delle classifiche (admin).
Verifica i fogli prodotti, che le righe coincidano con le classifiche in app,
e che un giocatore non possa scaricare il file.
"""

from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import load_workbook

from app.models.match import Match
from app.models.match_prediction import MatchPrediction
from app.models.round import Competition, Round, RoundStatus
from app.models.round_score import RoundScore
from app.models.season import Season, SeasonStatus
from app.models.user import User, UserRole
from app.services.auth import create_access_token, hash_password

pytestmark = pytest.mark.asyncio

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _auth(u: User) -> dict:
    return {"Authorization": f"Bearer {create_access_token({'sub': str(u.id), 'role': u.role})}"}


@pytest_asyncio.fixture
async def setup(db_session):
    """Stagione + admin + 2 player + 1 giornata completata con partite e schedine."""
    admin = User(
        username="admin", email="a@ex.com", password_hash=hash_password("x"), role=UserRole.admin
    )
    p1 = User(
        username="p1", email="p1@ex.com", password_hash=hash_password("x"), role=UserRole.player
    )
    p2 = User(
        username="p2", email="p2@ex.com", password_hash=hash_password("x"), role=UserRole.player
    )
    season = Season(name="2025-26", status=SeasonStatus.active)
    db_session.add_all([admin, p1, p2, season])
    await db_session.flush()

    kickoff = datetime(2026, 10, 21, 21, 0, tzinfo=timezone.utc)
    rnd = Round(
        season_id=season.id,
        name="G1",
        competition=Competition.serie_a,
        deadline=kickoff - timedelta(hours=2),
        status=RoundStatus.completed,
    )
    db_session.add(rnd)
    await db_session.flush()

    # Una partita col risultato esatto richiesto, una solo-segno
    m_exact = Match(
        round_id=rnd.id,
        home_team="Inter",
        away_team="Milan",
        competition=Competition.serie_a,
        requires_exact_score=True,
        kickoff=kickoff,
        actual_home_goals=2,
        actual_away_goals=1,
    )
    m_sign = Match(
        round_id=rnd.id,
        home_team="Roma",
        away_team="Lazio",
        competition=Competition.serie_a,
        requires_exact_score=False,
        kickoff=kickoff + timedelta(hours=2),
        actual_home_goals=0,
        actual_away_goals=0,
    )
    db_session.add_all([m_exact, m_sign])
    await db_session.flush()

    db_session.add_all(
        [
            # p1 azzecca il pieno (2-1) e il segno sul pareggio
            MatchPrediction(
                player_id=p1.id,
                match_id=m_exact.id,
                predicted_sign="1",
                predicted_home_goals=2,
                predicted_away_goals=1,
                points_earned=6,
            ),
            MatchPrediction(
                player_id=p1.id, match_id=m_sign.id, predicted_sign="X", points_earned=1
            ),
            # p2 sbaglia entrambe
            MatchPrediction(
                player_id=p2.id,
                match_id=m_exact.id,
                predicted_sign="2",
                predicted_home_goals=0,
                predicted_away_goals=1,
                points_earned=0,
            ),
            MatchPrediction(
                player_id=p2.id, match_id=m_sign.id, predicted_sign="1", points_earned=0
            ),
        ]
    )
    db_session.add_all(
        [
            RoundScore(
                player_id=p1.id,
                round_id=rnd.id,
                sign_points=2,
                exact_points=5,
                total_goals_points=0,
                weekend_bonus=6,
                total_round_points=13,
            ),
            RoundScore(
                player_id=p2.id,
                round_id=rnd.id,
                sign_points=0,
                exact_points=0,
                total_goals_points=0,
                weekend_bonus=4,
                total_round_points=4,
            ),
        ]
    )
    await db_session.commit()
    return {"admin": admin, "p1": p1, "p2": p2, "season": season, "round": rnd}


async def _workbook(client: AsyncClient, admin: User, **params):
    r = await client.get("/admin/export/leaderboards", headers=_auth(admin), params=params)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX_MEDIA_TYPE
    return r, load_workbook(BytesIO(r.content))


def _rows(ws) -> list[tuple]:
    """Righe dati (header escluso)."""
    return list(ws.iter_rows(min_row=2, values_only=True))


async def test_export_contiene_tutti_i_fogli(client: AsyncClient, setup):
    r, wb = await _workbook(client, setup["admin"])
    assert wb.sheetnames == ["Generale", "Speciali", "Giornate", "Schedine", "Info"]
    assert "totosport_classifiche_2025-26" in r.headers["content-disposition"]


async def test_foglio_generale_coincide_con_la_classifica_in_app(client: AsyncClient, setup):
    api = await client.get("/leaderboard", headers=_auth(setup["admin"]))
    expected = {row["username"]: row["total_points"] for row in api.json()}

    _, wb = await _workbook(client, setup["admin"])
    rows = _rows(wb["Generale"])
    assert {r[1]: r[2] for r in rows} == expected
    assert [r[0] for r in rows] == [1, 2]  # posizioni in ordine


async def test_foglio_speciali_impila_i_tre_ranking(client: AsyncClient, setup):
    _, wb = await _workbook(client, setup["admin"])
    rows = _rows(wb["Speciali"])
    assert {r[0] for r in rows} == {"Segni", "Pieni", "Pieni 5+ gol"}
    # p1 ha preso 2 segni e 1 pieno; p2 nessuno
    segni = {r[2]: r[3] for r in rows if r[0] == "Segni"}
    pieni = {r[2]: r[3] for r in rows if r[0] == "Pieni"}
    assert segni == {"p1": 2, "p2": 0}
    assert pieni == {"p1": 1, "p2": 0}


async def test_foglio_giornate_una_riga_per_giocatore(client: AsyncClient, setup):
    _, wb = await _workbook(client, setup["admin"])
    rows = _rows(wb["Giornate"])
    assert len(rows) == 2
    assert {r[3]: r[4] for r in rows} == {"p1": 13, "p2": 4}
    assert {r[0] for r in rows} == {"G1"}
    assert {r[1] for r in rows} == {"completed"}


async def test_foglio_schedine_riporta_segni_e_gol(client: AsyncClient, setup):
    _, wb = await _workbook(client, setup["admin"])
    rows = _rows(wb["Schedine"])
    assert len(rows) == 4

    inter = {r[2]: r for r in rows if r[1] == "Inter - Milan"}
    assert inter["p1"][3:9] == ("1", 2, 1, "sì", "2-1", 6)
    # Sulla partita solo-segno i gol restano vuoti
    roma = {r[2]: r for r in rows if r[1] == "Roma - Lazio"}
    assert roma["p1"][3:9] == ("X", None, None, "no", "0-0", 1)


async def test_schedine_escludibili(client: AsyncClient, setup):
    _, wb = await _workbook(client, setup["admin"], include_schedine=False)
    assert "Schedine" not in wb.sheetnames
    assert wb.sheetnames == ["Generale", "Speciali", "Giornate", "Info"]


async def test_un_giocatore_non_puo_scaricare(client: AsyncClient, setup):
    r = await client.get("/admin/export/leaderboards", headers=_auth(setup["p1"]))
    assert r.status_code == 403
