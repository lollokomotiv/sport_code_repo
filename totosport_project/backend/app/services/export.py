"""
Export delle classifiche in workbook Excel.

Due file distinti:
- `build_general_workbook`: classifica generale + classifiche speciali di stagione;
- `build_round_workbook`: punti e schedine di UNA giornata.

Le righe vengono SEMPRE prese dai service di classifica esistenti
(`general_leaderboard`, `special_rankings`, `round_leaderboard`): qui non c'è
nessuna query di scoring duplicata, così il file e l'app non possono divergere.
FastAPI-free come gli altri service: restituisce bytes, il router li serve.
"""

import re
import uuid
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.match import Match
from app.models.match_prediction import MatchPrediction
from app.models.round import Round
from app.models.season import Season
from app.models.user import User
from app.services.leaderboard import general_leaderboard, round_leaderboard, special_rankings

# Etichette dei tre ranking speciali, nell'ordine in cui compaiono in app
SPECIAL_LABELS: list[tuple[str, str]] = [
    ("segni", "Segni"),
    ("pieni", "Pieni"),
    ("pieni_5plus", "Pieni 5+ gol"),
]

HEADER_FONT = Font(bold=True)

# Una riga di foglio: celle eterogenee (testo, numeri, date, None)
Row = list[object]


def _write_sheet(ws: Worksheet, headers: list[str], rows: list[Row]) -> None:
    """Intestazione in grassetto, riga 1 congelata, autofiltro e colonne a misura."""
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    for idx, header in enumerate(headers, start=1):
        longest = max(
            [len(str(header))] + [len(str(r[idx - 1])) for r in rows if r[idx - 1] is not None]
            or [len(str(header))]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 40)


# ─── File 1: classifica generale di stagione ─────────────────────────────────


async def build_general_workbook(season: Season, db: AsyncSession) -> bytes:
    """Classifica generale + speciali della stagione. Nessun dato per giornata."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Generale"
    general = await general_leaderboard(season.id, db)
    _write_sheet(
        ws,
        [
            "Pos",
            "Giocatore",
            "Totale",
            "Segni",
            "Pieni",
            "Totale gol",
            "Bonus weekend",
            "Tabellone",
            "Bonus stagione",
        ],
        [
            [
                e.rank,
                e.username,
                e.total_points,
                e.sign_points,
                e.exact_points,
                e.total_goals_points,
                e.weekend_bonus_total,
                e.tabellone_points,
                e.season_bonus_total,
            ]
            for e in general
        ],
    )

    # Tre ranking speciali impilati, discriminante in colonna
    specials = await special_rankings(season.id, db)
    special_rows: list[Row] = []
    for key, label in SPECIAL_LABELS:
        for e in specials.get(key, []):
            special_rows.append([label, e["rank"], e["username"], e["count"], e["points"]])
    _write_sheet(
        wb.create_sheet("Speciali"),
        ["Classifica", "Pos", "Giocatore", "Azzeccate", "Punti"],
        special_rows,
    )

    _write_sheet(
        wb.create_sheet("Info"),
        ["Voce", "Valore"],
        [
            ["Stagione", season.name],
            ["Stato stagione", season.status.value],
            ["Generato il", datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)],
            ["Foglio Generale", "classifica generale, tutte le giornate con punteggio"],
            ["Foglio Speciali", "SOLO giornate in stato 'completed' (come in app)"],
            ["Dati per giornata", "scaricabili dal tab Dettaglio, un file per giornata"],
        ],
    )

    return _to_bytes(wb)


# ─── File 2: una singola giornata ────────────────────────────────────────────


async def _schedine_rows(round_id: uuid.UUID, db: AsyncSession) -> list[Row]:
    """Previsioni di una giornata, una riga per previsione."""
    result = await db.execute(
        select(MatchPrediction, Match, User.username)
        .join(Match, Match.id == MatchPrediction.match_id)
        .join(User, User.id == MatchPrediction.player_id)
        .where(Match.round_id == round_id)
        .order_by(Match.kickoff.nulls_last(), Match.home_team, User.username)
    )
    rows: list[Row] = []
    for mp, match, username in result.all():
        risultato = (
            f"{match.actual_home_goals}-{match.actual_away_goals}"
            if match.actual_home_goals is not None and match.actual_away_goals is not None
            else None
        )
        rows.append(
            [
                f"{match.home_team} - {match.away_team}",
                username,
                mp.predicted_sign,
                mp.predicted_home_goals,
                mp.predicted_away_goals,
                "sì" if match.requires_exact_score else "no",
                risultato,
                mp.points_earned,
                mp.submitted_at.replace(tzinfo=None) if mp.submitted_at else None,
            ]
        )
    return rows


async def build_round_workbook(rnd: Round, db: AsyncSession) -> bytes:
    """Punti e schedine della giornata indicata: due fogli, nient'altro."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Giornate"
    _write_sheet(
        ws,
        [
            "Giornata",
            "Stato",
            "Pos",
            "Giocatore",
            "Punti giornata",
            "Segni",
            "Pieni",
            "Totale gol",
            "Bonus weekend",
        ],
        [
            [
                rnd.name,
                rnd.status.value,
                e.rank,
                e.username,
                e.round_points,
                e.sign_points,
                e.exact_points,
                e.total_goals_points,
                e.weekend_bonus,
            ]
            for e in await round_leaderboard(rnd.id, db)
        ],
    )

    _write_sheet(
        wb.create_sheet("Schedine"),
        [
            "Partita",
            "Giocatore",
            "Segno",
            "Gol casa",
            "Gol ospiti",
            "Risultato esatto richiesto",
            "Risultato reale",
            "Punti",
            "Inviata il",
        ],
        await _schedine_rows(rnd.id, db),
    )

    return _to_bytes(wb)


# ─── Utilità ──────────────────────────────────────────────────────────────────


def _to_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _safe(part: str) -> str:
    """Ripulisce un pezzo di nome file dai caratteri problematici."""
    return re.sub(r'[\\/:*?"<>|]', "-", part).strip()


def general_filename(season: Season) -> str:
    return f"{_safe(season.name)}_Classifica_Generale.xlsx"


def round_filename(season: Season, rnd: Round) -> str:
    return f"{_safe(season.name)}_{_safe(rnd.name)}.xlsx"
